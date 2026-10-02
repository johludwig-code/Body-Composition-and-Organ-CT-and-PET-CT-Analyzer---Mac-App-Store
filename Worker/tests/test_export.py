"""The export of plan §11 on a synthetic project (tests/export_project.py).

Golden files live in tests/golden/export/. They are CSV because CSV is
byte-comparable; the workbook is read back with openpyxl and compared against
the same tables. Regenerate after an intended change with
`BCOA_UPDATE_GOLDEN=1 pytest tests/test_export.py` and review the diff.
"""

from __future__ import annotations

import csv
import dataclasses
import hashlib
import io
import json
import os
import re
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

import export_project
from bcoa_worker.channel import ProtocolChannel
from bcoa_worker.errors import JobFailure
from bcoa_worker.export import data as export_data
from bcoa_worker.export.options import ExportOptionError, parse_options
from bcoa_worker.export.tables import SHEET_ORDER, Column, Sheet, build
from bcoa_worker.export.writers import write_csv, write_xlsx
from bcoa_worker.jobs import export as export_job
from bcoa_worker.metrics import METRIC_COLUMNS
from bcoa_worker.protocol import Job
from conftest import WORKER_ROOT

GOLDEN = WORKER_ROOT / "tests" / "golden" / "export"
CREATED = "2026-10-02T14:00:00Z"


@pytest.fixture(scope="module")
def project(tmp_path_factory: pytest.TempPathFactory) -> export_data.ProjectData:
    folder = tmp_path_factory.mktemp("project")
    return export_data.load(export_project.create(folder), export_project.RUN)


def _export(project: export_data.ProjectData, **options: object):
    return build(project, parse_options(options), created=CREATED)


def _rows(sheet: Sheet) -> list[dict[str, object]]:
    names = [c.name for c in sheet.columns]
    return [dict(zip(names, row, strict=True)) for row in sheet.rows]


def _by_pseudonym(sheet: Sheet) -> dict[str, dict[str, object]]:
    return {str(r["pseudonym"]): r for r in _rows(sheet)}


# -- options -------------------------------------------------------------------


def test_defaults_are_the_plans_defaults() -> None:
    options = parse_options({})
    assert (options.format, options.layout, options.row_level) == ("xlsx", "wide", "patient")
    assert (options.multiple_studies, options.multiple_series) == ("suffix", "primary")
    assert (options.qc, options.identifiers, options.dates) == (
        "exclude",
        "pseudonym",
        "days_since_first",
    )
    assert options.methods_text is True and options.masks is False
    assert options.metrics == METRIC_COLUMNS


@pytest.mark.parametrize(
    "raw",
    [
        {"layout": "tall"},
        {"format": "xls"},
        {"colour": "red"},
        {"metrics": ["hu_mode"]},
        {"models": []},
        {"labels": "liver"},
        {"reproducibility_package": True},
    ],
)
def test_options_the_dialog_cannot_send_are_refused(raw: dict[str, object]) -> None:
    with pytest.raises(ExportOptionError):
        parse_options(raw)


def test_metrics_come_out_in_plan_order_whatever_order_was_sent() -> None:
    assert parse_options({"metrics": ["hu_max", "volume_ml"]}).metrics == ("volume_ml", "hu_max")


# -- reading the database --------------------------------------------------------


def test_loader_reads_without_changing_the_file(tmp_path: Path) -> None:
    database = export_project.create(tmp_path)
    before = hashlib.sha256(database.read_bytes()).hexdigest()
    project = export_data.load(database, export_project.RUN)
    assert hashlib.sha256(database.read_bytes()).hexdigest() == before
    assert [m.name for m in project.run.models] == ["clin_ct_organs", "clin_ct_body_composition"]
    dates = {s.study_key: s.study_date for s in project.studies}
    # ISO and DICOM dates both parse; a malformed one is missing, not guessed.
    assert str(dates["st_b"]) == "2020-01-10" and str(dates["st_a"]) == "2021-03-05"
    assert dates["st_e"] is None


def test_unknown_run_is_a_lookup_error(tmp_path: Path) -> None:
    with pytest.raises(LookupError):
        export_data.load(export_project.create(tmp_path), "R_missing")


# -- wide layout -------------------------------------------------------------------


def test_wide_has_one_row_per_patient_sorted_by_pseudonym(project) -> None:
    results = _export(project).sheet("results")
    assert [r[0] for r in results.rows] == ["P0001", "P0002", "P0003", "P0004"]


def test_wide_header_starts_as_in_the_plan(project) -> None:
    names = [c.name for c in _export(project).sheet("results").columns]
    assert names[:12] == [
        "pseudonym",
        "n_studies",
        "sex",
        "age_first_study",
        "day__t1",
        "series_description__t1",
        "slice_thickness_mm__t1",
        "kernel__t1",
        "device__t1",
        "qc_status__t1",
        "run_id__t1",
        "organs__status__t1",
    ]
    assert names[12] == "organs__liver__voxel_count__t1"
    # 4 fixed + 2 timepoints × (7 per-timepoint + 2 models × (1 status + 3 labels × 10)).
    assert len(names) == 4 + 2 * (7 + 2 * (1 + 3 * 10))
    assert len(names) == len(set(names))


def test_timepoints_follow_the_study_date_not_the_key(project) -> None:
    row = _by_pseudonym(_export(project).sheet("results"))["P0001"]
    # st_b (2020) sorts after st_a (2021) by key but is t1 by date; the
    # scout-only study of 2019 has no series in the run and is no timepoint.
    assert row["n_studies"] == 2
    assert row["day__t1"] == 0 and row["day__t2"] == 420
    assert row["organs__liver__voxel_count__t1"] == 1010
    assert row["organs__liver__voxel_count__t2"] == 3010
    assert row["organs__liver__volume_ml__t2"] == 3010 / 128


def test_a_patient_with_one_study_has_empty_t2_cells(project) -> None:
    row = _by_pseudonym(_export(project).sheet("results"))["P0004"]
    assert row["organs__status__t2"] is None
    assert row["organs__liver__volume_ml__t2"] is None


def test_rejected_series_is_excluded_with_empty_cells_not_zeros(project) -> None:
    row = _by_pseudonym(_export(project).sheet("results"))["P0002"]
    assert row["qc_status__t1"] == "rejected"
    assert row["organs__status__t1"] == "excluded"
    assert all(row[f"organs__liver__{m}__t1"] is None for m in METRIC_COLUMNS)


def test_rejected_series_is_kept_with_its_status_when_asked(project) -> None:
    row = _by_pseudonym(_export(project, qc="include").sheet("results"))["P0002"]
    assert row["qc_status__t1"] == "rejected"
    assert row["organs__status__t1"] == "ok"
    assert row["organs__liver__voxel_count__t1"] == 7010


def test_failed_and_not_run_models_say_so(project) -> None:
    rows = _by_pseudonym(_export(project).sheet("results"))
    assert rows["P0003"]["organs__status__t1"] == "failed"
    assert rows["P0004"]["organs__status__t1"] == "ok"
    assert rows["P0004"]["body_composition__status__t1"] == "not_run"


def test_excluded_label_is_empty_and_its_neighbours_are_not(project) -> None:
    row = _by_pseudonym(_export(project).sheet("results"))["P0004"]
    assert row["organs__spleen__volume_ml__t1"] is None
    assert row["organs__liver__volume_ml__t1"] == 9010 / 128
    assert row["organs__liver__touches_border__t1"] is True


def test_first_and_last_drop_the_suffix(project) -> None:
    last = _export(project, multiple_studies="last").sheet("results")
    names = [c.name for c in last.columns]
    assert "organs__liver__volume_ml" in names
    assert not any(re.search(r"__t\d", n) for n in names)
    row = _by_pseudonym(last)["P0001"]
    assert row["organs__liver__voxel_count"] == 3010
    # Still the number of studies the patient has.
    assert row["n_studies"] == 2
    first = _by_pseudonym(_export(project, multiple_studies="first").sheet("results"))
    assert first["P0001"]["organs__liver__voxel_count"] == 1010


def test_all_series_get_their_own_suffix(project) -> None:
    results = _export(project, multiple_series="all").sheet("results")
    row = _by_pseudonym(results)["P0001"]
    assert row["organs__liver__voxel_count__t2__s1"] == 3010
    assert row["organs__liver__voxel_count__t2__s2"] == 5010
    # No patient has a second series at t1, so there is no column for one.
    assert "organs__liver__voxel_count__t1__s2" not in row


def test_one_row_per_study_and_per_series(project) -> None:
    per_study = _export(project, row_level="study").sheet("results")
    assert [(r[0], r[4]) for r in per_study.rows] == [
        ("P0001", 1),
        ("P0001", 2),
        ("P0002", 1),
        ("P0003", 1),
        ("P0004", 1),
    ]
    per_series = _export(project, row_level="series", multiple_series="all").sheet("results")
    assert [(r[0], r[4], r[5]) for r in per_series.rows] == [
        ("P0001", 1, 1),
        ("P0001", 2, 1),
        ("P0001", 2, 2),
        ("P0002", 1, 1),
        ("P0003", 1, 1),
        ("P0004", 1, 1),
    ]


def test_subsets_of_patients_models_labels_and_metrics(project) -> None:
    results = _export(
        project,
        patients=["P0004", "P0001"],
        models=["clin_ct_organs"],
        labels=["Liver"],
        metrics=["volume_ml"],
        multiple_studies="first",
    ).sheet("results")
    assert [c.name for c in results.columns][-2:] == ["organs__status", "organs__liver__volume_ml"]
    assert [r[0] for r in results.rows] == ["P0001", "P0004"]


def test_a_model_that_is_not_in_the_run_is_refused(project) -> None:
    with pytest.raises(ValueError, match="not in run"):
        _export(project, models=["clin_ct_ribs"])


@pytest.mark.parametrize(
    ("dates", "column", "value"),
    [
        ("days_since_first", "day__t2", 420),
        ("year", "year__t2", 2021),
        ("iso", "study_date__t2", "2021-03-05"),
    ],
)
def test_date_options(project, dates: str, column: str, value: object) -> None:
    assert _by_pseudonym(_export(project, dates=dates).sheet("results"))["P0001"][column] == value


def test_patient_ids_come_with_a_warning(project) -> None:
    export = _export(project, identifiers="both")
    results = export.sheet("results")
    assert [c.name for c in results.columns][:2] == ["pseudonym", "patient_id"]
    assert [r[1] for r in results.rows] == ["00012345", "1E5", None, "4"]
    assert any("PatientID" in w for w in export.warnings)
    assert any("removed" in w for w in export.warnings)
    assert _export(project).warnings == []


# -- long layout ----------------------------------------------------------------------


def test_long_has_one_row_per_series_model_and_label(project) -> None:
    rows = _rows(_export(project, layout="long").sheet("results"))
    # P0001: 2 timepoints × (3 + 3); P0002 rejected; P0003 failed; P0004: organs
    # without the excluded spleen.
    assert len(rows) == 14
    assert {r["pseudonym"] for r in rows} == {"P0001", "P0004"}
    assert [r["label"] for r in rows if r["pseudonym"] == "P0004"] == ["liver", "kidney_left"]
    liver = next(r for r in rows if r["pseudonym"] == "P0004" and r["label"] == "liver")
    assert liver["flags"] == "truncated" and liver["touches_border"] is True


def test_long_with_qc_included_marks_rejected_and_excluded(project) -> None:
    rows = _rows(_export(project, layout="long", qc="include").sheet("results"))
    assert len(rows) == 14 + 6 + 1
    assert {r["qc_status"] for r in rows if r["pseudonym"] == "P0002"} == {"rejected"}
    spleen = next(r for r in rows if r["pseudonym"] == "P0004" and r["label"] == "spleen")
    assert spleen["qc_status"] == "label_excluded" and spleen["volume_ml"] == 9020 / 128


# -- the other sheets ------------------------------------------------------------------


def test_workbook_has_the_plans_sheets_in_order(project) -> None:
    assert tuple(s.name for s in _export(project).sheets) == SHEET_ORDER


def test_data_dictionary_describes_every_column(project) -> None:
    export = _export(project, layout="long")
    described = {(r[0], r[1]) for r in export.sheet("data_dictionary").rows}
    for sheet in export.sheets:
        if sheet.name != "data_dictionary":
            assert {(sheet.name, c.name) for c in sheet.columns} <= described


def test_skipped_names_the_reason(project) -> None:
    rows = _rows(_export(project).sheet("skipped"))
    assert [(r["pseudonym"], r["model"], r["reason"]) for r in rows] == [
        ("P0002", None, "excluded by QC"),
        ("P0003", None, "failed: dicom_unreadable"),
        ("P0004", "body_composition", "not run"),
    ]


def test_qc_sheet_has_the_latest_decision_and_the_label_exclusion(project) -> None:
    rows = _rows(_export(project).sheet("qc"))
    p2 = [r for r in rows if r["pseudonym"] == "P0002"]
    assert [(r["status"], r["comment"]) for r in p2] == [("rejected", "motion artefact")]
    p4 = [(r["label"], r["status"], r["flags"]) for r in rows if r["pseudonym"] == "P0004"]
    assert ("spleen", "label_excluded", None) in p4
    assert ("liver", None, "truncated") in p4


def test_series_sheet_carries_runtime_and_primary_flag(project) -> None:
    rows = _rows(_export(project, multiple_series="all").sheet("series"))
    p1 = [
        (r["timepoint"], r["series"], r["primary"], r["runtime_s"])
        for r in rows
        if r["pseudonym"] == "P0001"
    ]
    assert p1 == [(1, 1, True, 90.0), (2, 1, True, 60.0), (2, 2, False, 60.0)]


def test_methods_text_and_provenance(project) -> None:
    export = _export(project)
    assert export.methods_text == (
        "Segmentations were generated with MOOSE v3.2.2 (Shiyam Sundar et al., J Nucl Med "
        "2022), based on nnU-Net (Isensee et al., Nat Methods 2021), using Body Composition "
        "and Organ CT and PET-CT Analyzer v0.1.0 on Apple M1 Pro (PyTorch mps). 3 of 4 series "
        "underwent visual quality control; 1 series were excluded."
    )
    provenance = dict(export.sheet("provenance").rows)
    assert provenance["intended_use"] == "For research use only. Not for clinical use."
    assert provenance["model.clin_ct_organs.sha256"] == "a" * 64
    assert "CC BY 4.0" in str(provenance["model_attribution"])
    assert json.loads(str(provenance["export_options"]))["layout"] == "wide"


def _qc_sentence(project: export_data.ProjectData) -> str:
    return _export(project).methods_text.split("(PyTorch mps). ")[1]


def test_methods_text_claims_no_review_that_was_not_recorded(project) -> None:
    # A real case exported straight after segmentation said "Results underwent
    # visual quality control" although nobody had looked at it.
    unreviewed = dataclasses.replace(project, qc=())
    assert _qc_sentence(unreviewed) == "No visual quality control was recorded for these results."


def test_methods_text_uses_the_plan_sentence_once_every_series_was_reviewed(project) -> None:
    seen = [
        export_data.QCEntry(key, None, None, "accepted", "JL", "2026-10-02T14:00:00Z", None)
        for key in ("se2", "se6")
    ]
    reviewed = dataclasses.replace(project, qc=(*project.qc, *seen))
    assert _qc_sentence(reviewed) == (
        "Results underwent visual quality control; 1 series were excluded."
    )


# -- writers ---------------------------------------------------------------------------


def test_workbook_reads_back_with_the_right_types(project, tmp_path: Path) -> None:
    export = _export(project, identifiers="both")
    path = tmp_path / "export.xlsx"
    write_xlsx(export.sheets, path, created=datetime(2026, 10, 2, 14, 0, 0))
    workbook = openpyxl.load_workbook(path)
    assert tuple(workbook.sheetnames) == SHEET_ORDER
    sheet = workbook["results"]
    header = [c.value for c in sheet[1]]
    assert header == [c.name for c in export.sheet("results").columns]
    rows = {
        r[0]: dict(zip(header, r, strict=True))
        for r in sheet.iter_rows(min_row=2, values_only=True)
    }
    # Leading zeros and the scientific-notation trap survive as text.
    assert rows["P0001"]["patient_id"] == "00012345"
    assert rows["P0002"]["patient_id"] == "1E5"
    assert rows["P0001"]["organs__liver__volume_ml__t2"] == 3010 / 128
    assert rows["P0001"]["organs__liver__voxel_count__t2"] == 3010
    assert rows["P0004"]["organs__liver__touches_border__t1"] is True
    # Missing is an empty cell, never 0 or "NA".
    assert rows["P0002"]["organs__liver__volume_ml__t1"] is None
    assert rows["P0004"]["series_description__t1"] == '=HYPERLINK("http://example.invalid")'
    formula_cell = next(c for c in sheet["A"][1:] if c.value == "P0004").offset(
        column=header.index("series_description__t1")
    )
    assert formula_cell.data_type == "s"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None
    # Every value matches the table it came from.
    table = export.sheet("results")
    for got, want in zip(sheet.iter_rows(min_row=2, values_only=True), table.rows, strict=True):
        assert list(got) == want


def test_same_export_gives_the_same_workbook(project, tmp_path: Path) -> None:
    export = _export(project)
    created = datetime(2026, 10, 2, 14, 0, 0)
    write_xlsx(export.sheets, tmp_path / "a.xlsx", created=created)
    write_xlsx(export.sheets, tmp_path / "b.xlsx", created=created)
    assert (tmp_path / "a.xlsx").read_bytes() == (tmp_path / "b.xlsx").read_bytes()


def test_csv_for_excel_uses_semicolon_decimal_comma_bom_and_defuses_formulas(
    project, tmp_path: Path
) -> None:
    path = tmp_path / "excel.csv"
    write_csv(_export(project, layout="long").sheet("results"), path, "csv_excel")
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    header, body = rows[0], rows[1:]
    p4 = [dict(zip(header, r, strict=True)) for r in body if r[0] == "P0004"]
    assert p4[0]["series_description"] == '\'=HYPERLINK("http://example.invalid")'
    assert p4[0]["volume_ml"] == str(9010 / 128).replace(".", ",")
    assert p4[0]["touches_border"] == "TRUE"


def test_csv_for_r_and_python_is_verbatim_and_reads_back(project, tmp_path: Path) -> None:
    pandas = pytest.importorskip("pandas")
    path = tmp_path / "plain.csv"
    write_csv(_export(project, identifiers="both").sheet("results"), path, "csv")
    assert not path.read_bytes().startswith(b"\xef\xbb\xbf")
    frame = pandas.read_csv(path, dtype={"patient_id": str})
    assert list(frame["patient_id"].fillna("")) == ["00012345", "1E5", "", "4"]
    assert frame.loc[0, "organs__liver__volume_ml__t2"] == 3010 / 128
    assert pandas.isna(frame.loc[1, "organs__liver__volume_ml__t1"])
    assert frame.loc[3, "series_description__t1"] == '=HYPERLINK("http://example.invalid")'


def test_non_finite_numbers_become_empty_cells(tmp_path: Path) -> None:
    sheet = Sheet(
        "results",
        [Column("id", "text"), Column("x", "number")],
        [["a", float("nan")], ["b", float("inf")], ["c", 1.5]],
    )
    write_csv(sheet, tmp_path / "x.csv", "csv")
    assert (tmp_path / "x.csv").read_text().splitlines() == ["id,x", "a,", "b,", "c,1.5"]


def test_too_many_columns_fail_with_advice(tmp_path: Path) -> None:
    sheet = Sheet("results", [Column(f"c{i}", "number") for i in range(16_385)], [])
    with pytest.raises(ValueError, match="long layout"):
        write_xlsx([sheet], tmp_path / "wide.xlsx", created=datetime(2026, 1, 1))


# -- golden files ------------------------------------------------------------------------

GOLDEN_CASES = {
    "wide_patient": {},
    "wide_all_series": {"multiple_series": "all"},
    "wide_last_year": {"multiple_studies": "last", "dates": "year"},
    "wide_per_series": {"row_level": "series", "multiple_series": "all"},
    "long": {"layout": "long"},
    "long_qc_included": {"layout": "long", "qc": "include"},
}


@pytest.mark.parametrize("case", sorted(GOLDEN_CASES))
def test_golden_files(project, tmp_path: Path, case: str) -> None:
    export = _export(project, **GOLDEN_CASES[case])
    for dialect in ("csv", "csv_excel"):
        name = f"{case}.{'excel.' if dialect == 'csv_excel' else ''}csv"
        produced = tmp_path / name
        write_csv(export.sheet("results"), produced, dialect)
        golden = GOLDEN / name
        if os.environ.get("BCOA_UPDATE_GOLDEN"):
            golden.parent.mkdir(parents=True, exist_ok=True)
            golden.write_bytes(produced.read_bytes())
        assert produced.read_bytes() == golden.read_bytes(), f"{name} differs from its golden file"


# -- the job --------------------------------------------------------------------------------


def _job(project_dir: Path, payload: dict[str, object]) -> Job:
    return Job(
        "j_export",
        "export",
        project_dir,
        project_dir / "logs" / "j_export.log",
        project_dir,
        payload,
    )


def _events(buffer: io.StringIO) -> list[dict[str, object]]:
    return [json.loads(line) for line in buffer.getvalue().splitlines()]


def test_job_writes_the_workbook_methods_and_masks(tmp_path: Path) -> None:
    export_project.create(tmp_path)
    mask = tmp_path / "work" / "se1" / "labels" / "clin_ct_organs.nii.gz"
    mask.parent.mkdir(parents=True)
    mask.write_bytes(b"labelmap")
    buffer = io.StringIO()
    export_job.run(
        _job(
            tmp_path, {"run_id": export_project.RUN, "stem": "cohort", "options": {"masks": True}}
        ),
        ProtocolChannel(buffer),
    )
    events = _events(buffer)
    artifacts = {e["path"]: e["sha256"] for e in events if e["type"] == "artifact"}
    assert "exports/cohort.xlsx" in artifacts
    assert "exports/cohort.methods.txt" in artifacts
    assert "exports/cohort_masks/P0001_t1_s1_organs.nii.gz" in artifacts
    for path, digest in artifacts.items():
        assert hashlib.sha256((tmp_path / path).read_bytes()).hexdigest() == digest
    result = next(e for e in events if e["type"] == "result")["payload"]
    assert result["rows"] == 4 and result["excluded_series"] == 1
    # Masks that are not in the project folder are reported, not invented.
    assert any(e["type"] == "log" and "missing" in e["message"] for e in events)
    assert not list((tmp_path / "exports").glob(".*partial"))


def test_job_writes_one_csv_per_sheet(tmp_path: Path) -> None:
    export_project.create(tmp_path)
    export_job.run(
        _job(
            tmp_path, {"run_id": export_project.RUN, "stem": "cohort", "options": {"format": "csv"}}
        ),
        ProtocolChannel(io.StringIO()),
    )
    names = sorted(p.name for p in (tmp_path / "exports").iterdir())
    assert names == sorted(
        ["cohort.csv", "cohort.methods.txt"]
        + [f"cohort.{s}.csv" for s in SHEET_ORDER if s != "results"]
    )


@pytest.mark.parametrize("stem", ["../escape", "a/b", "", ".hidden", "x."])
def test_job_refuses_names_that_are_not_plain_file_names(tmp_path: Path, stem: str) -> None:
    export_project.create(tmp_path)
    with pytest.raises(JobFailure) as failure:
        export_job.run(
            _job(tmp_path, {"run_id": export_project.RUN, "stem": stem}),
            ProtocolChannel(io.StringIO()),
        )
    assert failure.value.code == "bad_export_name"


def test_job_reports_an_unknown_run(tmp_path: Path) -> None:
    export_project.create(tmp_path)
    with pytest.raises(JobFailure) as failure:
        export_job.run(
            _job(tmp_path, {"run_id": "R_x", "stem": "x"}), ProtocolChannel(io.StringIO())
        )
    assert failure.value.code == "unknown_run"


def test_the_protocol_fixture_is_a_valid_export_job() -> None:
    fixture = WORKER_ROOT.parent / "Protocol" / "fixtures" / "jobs" / "export.json"
    payload = json.loads(fixture.read_text())["payload"]
    assert export_job._STEM.match(payload["stem"])
    assert parse_options(payload["options"]).layout == "wide"
